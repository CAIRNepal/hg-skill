#!/usr/bin/env python3
"""Ingestion engine for tabular (CSV, TSV, XLSX) and JSON cultural heritage records.

Isolates structured metadata (IDs, coordinates, types, admin levels)
from unstructured narrative prose (descriptions, inscriptions, histories)
to prepare records for LinkML Container extraction.
"""
import csv
import json
from pathlib import Path
from typing import Any, Dict, List, Set, Tuple

import openpyxl

NULL_STRINGS = {"", "null", "none", "n/a", "na", "undefined", "unknown", "-", "--"}

NARRATIVE_KEYWORDS = {
    "desc", "description", "history", "historical", "narrative", "note", "notes",
    "inscription", "inscriptions", "text", "transcription", "ritual", "rituals",
    "summary", "detail", "details", "oral_history", "legend", "mythology"
}


def clean_record(raw: Dict[str, Any]) -> Dict[str, Any]:
    """Clean empty values, whitespace, and null-like strings from a record."""
    cleaned = {}
    for k, v in raw.items():
        if v is None:
            continue
        if isinstance(v, str):
            val = v.strip()
            if val.lower() in NULL_STRINGS:
                continue
            cleaned[str(k).strip()] = val
        else:
            cleaned[str(k).strip()] = v
    return cleaned


def categorize_fields(fields: List[str]) -> Tuple[Set[str], Set[str]]:
    """Partition fields into (structured_fields, narrative_fields) using token matching."""
    structured = set()
    narrative = set()

    for f in fields:
        f_norm = f.lower().replace("-", "_").replace(" ", "_")
        tokens = set(filter(None, f_norm.split("_")))

        if tokens & NARRATIVE_KEYWORDS:
            narrative.add(f)
        else:
            structured.add(f)

    return structured, narrative


def partition_record(rec: Dict[str, Any], structured_fields: Set[str], narrative_fields: Set[str]) -> Dict[str, Any]:
    """Partition a single cleaned record into structured vs narrative components."""
    rec_id = None
    for id_key in ("id", "monument_id", "accession_no", "resource_id", "code"):
        if id_key in rec:
            rec_id = str(rec[id_key])
            break

    name = None
    for name_key in ("name", "monument_name", "title", "label"):
        if name_key in rec:
            name = str(rec[name_key])
            break

    structured = {}
    narrative = {}

    for k, v in rec.items():
        if k in narrative_fields or (isinstance(v, str) and len(v) > 200):
            narrative[k] = v
        else:
            structured[k] = v

    out = {
        "structured": structured,
        "narrative": narrative
    }
    if rec_id:
        out["id"] = rec_id
    if name:
        out["name"] = name
    return out


def ingest_csv(file_path: Path, delimiter: str = None) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Ingest CSV or TSV file into partitioned records with encoding fallback."""
    p = Path(file_path)
    if delimiter is None:
        delimiter = "\t" if p.suffix.lower() == ".tsv" else ","

    records = []
    fields = []
    used_encoding = "utf-8"

    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with open(p, "r", encoding=enc, newline="") as f:
                reader = csv.DictReader(f, delimiter=delimiter)
                fields = reader.fieldnames or []
                structured_fields, narrative_fields = categorize_fields(fields)
                temp_records = []
                for row in reader:
                    cleaned = clean_record(row)
                    if cleaned:
                        temp_records.append(partition_record(cleaned, structured_fields, narrative_fields))
                records = temp_records
                used_encoding = enc
                break
        except UnicodeDecodeError:
            continue

    metadata = {
        "source_file": str(p.name),
        "format": "tsv" if delimiter == "\t" else "csv",
        "record_count": len(records),
        "fields": fields
    }
    return records, metadata


def ingest_excel(file_path: Path, sheet_name: str = None) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Ingest Excel (.xlsx) file into partitioned records."""
    p = Path(file_path)
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], {"source_file": str(p.name), "format": "xlsx", "record_count": 0}

    header_row = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    structured_fields, narrative_fields = categorize_fields(header_row)

    records = []
    for r in rows[1:]:
        row_dict = {}
        for col_name, val in zip(header_row, r):
            if val is not None:
                row_dict[col_name] = val
        cleaned = clean_record(row_dict)
        if cleaned:
            records.append(partition_record(cleaned, structured_fields, narrative_fields))

    metadata = {
        "source_file": str(p.name),
        "format": "xlsx",
        "sheet_name": ws.title,
        "record_count": len(records),
        "fields": header_row
    }
    return records, metadata


def ingest_json(file_path: Path) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Ingest JSON file (array of records or dictionary) into partitioned records."""
    p = Path(file_path)
    data = json.load(open(p, encoding="utf-8"))

    raw_list = []
    if isinstance(data, list):
        raw_list = data
    elif isinstance(data, dict):
        # Look for common wrapper array keys
        for k in ("records", "items", "monuments", "data", "results"):
            if k in data and isinstance(data[k], list):
                raw_list = data[k]
                break
        if not raw_list:
            raw_list = [data]

    # Collect fields across records
    fields = set()
    for item in raw_list:
        if isinstance(item, dict):
            fields.update(item.keys())

    structured_fields, narrative_fields = categorize_fields(list(fields))

    records = []
    for item in raw_list:
        if isinstance(item, dict):
            cleaned = clean_record(item)
            if cleaned:
                records.append(partition_record(cleaned, structured_fields, narrative_fields))

    metadata = {
        "source_file": str(p.name),
        "format": "json",
        "record_count": len(records),
        "fields": list(fields)
    }
    return records, metadata
